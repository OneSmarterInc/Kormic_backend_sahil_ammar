import json
import tempfile
import uuid
from pathlib import Path
from unittest import mock

from django.test import TestCase, override_settings
from django.contrib.auth.models import User
from langchain_core.messages import AIMessage, HumanMessage
from langgraph.checkpoint.memory import InMemorySaver

from accounts.models import Account
from django_api.models import (StudentProfile, AgentChangeProposal, UniversityKnowledgeEntry,
    UniversityInterestEvent, ChatMessage, ChatAttachment, ResumeUpload, LinkedInAnalysis, KnowledgeIndexWork)
from django_api.services import profile_row_to_dict
from universities.models import University
from pure_multi_agent import change_proposals as changes
from pure_multi_agent import document_evidence as documents


@override_settings(UNIVERSITY_VECTOR_SEARCH=False, AGENT_DISTRIBUTED_LIMITS=False)
class ChangeTests(TestCase):
    def setUp(self):
        self.student = StudentProfile.objects.create(name='Asha', gpa=3.4, gpa_scale='4.0')
        self.ctx = self.context()

    def context(self, message='My GPA is 3.8', turn='one'):
        self.student.refresh_from_db()
        profile = profile_row_to_dict(self.student)
        return {'canonical_student_id': str(self.student.uuid), 'student_profile': profile,
            'profile_baseline': dict(profile), 'current_message': message, 'turn_id': turn}

    def proposal(self, **values):
        return changes.update_student(self.ctx, values)['confirmation_required']

    def test_missing_same_and_conflicting_values_in_one_request(self):
        result = changes.update_student(self.ctx, {'gpa': 3.8, 'gpa_scale': 4.0, 'major': 'Computer Science'})
        self.student.refresh_from_db()
        self.assertEqual(self.student.major, 'Computer Science')
        self.assertEqual(self.student.gpa, 3.4)
        self.assertEqual(result['unchanged_fields'], ['gpa_scale'])
        self.assertEqual(result['confirmation_required']['before'], {'gpa': 3.4})

    def test_same_turn_cannot_confirm(self):
        proposal = self.proposal(gpa=3.8)
        with self.assertRaisesMessage(ValueError, 'later user turn'):
            changes.resolve(self.ctx, proposal['id'], 'approve', self.ctx['current_message'])

    def test_later_approval_persists_exact_values_idempotently(self):
        proposal = self.proposal(gpa=3.8)
        ctx = self.context('Yes, save it', 'two')
        for _ in range(2):
            self.assertEqual(changes.resolve(ctx, proposal['id'], 'approve', ctx['current_message'])['status'], 'applied')
        self.student.refresh_from_db()
        self.assertEqual(self.student.gpa, 3.8)
        self.assertEqual(AgentChangeProposal.objects.count(), 1)

    def test_reject_keeps_database_and_uses_temporary_assumption(self):
        proposal = self.proposal(gpa=3.8)
        ctx = self.context('No, do not save', 'two')
        changes.resolve(ctx, proposal['id'], 'reject', ctx['current_message'])
        effective, assumptions = changes.effective_profile(ctx)
        self.student.refresh_from_db()
        self.assertEqual(self.student.gpa, 3.4)
        self.assertEqual(ctx['student_profile']['gpa'], 3.4)
        self.assertEqual(effective['gpa'], 3.8)
        self.assertEqual(assumptions, {'gpa': 3.8})
        changes.clear_conversation(student_id=str(self.student.uuid))
        self.assertEqual(changes.effective_profile(ctx)[1], {})

    def test_cancel_is_not_an_assumption(self):
        proposal = self.proposal(gpa=3.8)
        ctx = self.context('Forget that', 'two')
        changes.resolve(ctx, proposal['id'], 'cancel', ctx['current_message'])
        self.assertFalse(changes.conversation_state(ctx)['conversation_assumptions'])

    def test_stale_approval_does_not_overwrite_external_edit(self):
        proposal = self.proposal(gpa=3.8)
        StudentProfile.objects.filter(pk=self.student.pk).update(gpa=3.6)
        ctx = self.context('yes', 'two')
        self.assertEqual(changes.resolve(ctx, proposal['id'], 'approve', 'yes')['status'], 'stale')
        self.student.refresh_from_db()
        self.assertEqual(self.student.gpa, 3.6)

    def test_foreign_proposal_and_fabricated_consent_are_rejected(self):
        proposal = self.proposal(gpa=3.8)
        other = StudentProfile.objects.create(name='Other')
        ctx = {**self.ctx, 'canonical_student_id': str(other.uuid), 'turn_id': 'two', 'current_message': 'yes'}
        with self.assertRaisesMessage(ValueError, 'not found'):
            changes.resolve(ctx, proposal['id'], 'approve', 'yes')
        with self.assertRaisesMessage(ValueError, 'entire current'):
            changes.resolve({**self.ctx, 'turn_id': 'two'}, proposal['id'], 'approve', 'yes')

    def test_invalid_and_unknown_fields_do_not_write(self):
        for values in ({'gpa': 5}, {'ielts': 15}, {'budget': -1}, {'gpa': float('nan')}, {'email': 'x@y.z'}):
            with self.subTest(values=values), self.assertRaises(ValueError):
                changes.update_student(self.ctx, values)
        self.assertEqual(AgentChangeProposal.objects.count(), 0)

    def test_document_turn_cannot_bypass_confirmation_with_profile_tool(self):
        self.ctx['chat_attachments'] = [{'id': 1}]
        with self.assertRaisesMessage(ValueError, 'Document updates always need confirmation'):
            changes.update_student(self.ctx, {'major': 'CS'})

    def test_revised_profile_proposal_supersedes_old_values(self):
        old = self.proposal(gpa=3.8)
        self.ctx = self.context('Actually 3.9', 'two')
        new = self.proposal(gpa=3.9)
        self.assertEqual(AgentChangeProposal.objects.get(pk=old['id']).status, 'superseded')
        self.assertEqual(changes.conversation_state(self.ctx)['pending_changes'][0]['id'], new['id'])

    def test_history_refreshes_proposal_status(self):
        proposal = self.proposal(gpa=3.8)
        message = ChatMessage.objects.create(student_id=str(self.student.uuid), channel='agent', sender='assistant', content='Confirm?', meta={'change_proposals': [proposal]})
        changes.resolve(self.context('yes', 'two'), proposal['id'], 'approve', 'yes')
        changes.refresh_message_metadata([message], student_id=str(self.student.uuid))
        self.assertEqual(message.meta['change_proposals'][0]['status'], 'applied')


class DocumentTests(TestCase):
    context = ChangeTests.context
    def setUp(self):
        ChangeTests.setUp(self)
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        override = override_settings(MEDIA_ROOT=self.temp.name)
        override.enable()
        self.addCleanup(override.disable)
        self.path = Path(self.temp.name) / 'resume.txt'
        self.path.write_text('Asha Raman\nGPA 3.8/4.0\nPython and SQL', encoding='utf-8')
        msg = ChatMessage.objects.create(channel='agent', student_id=str(self.student.uuid), sender='user', content='My resume')
        self.attachment = ChatAttachment.objects.create(message=msg, file_path=str(self.path), original_filename='resume.txt', content_type='text/plain')

    def document_proposal(self, kind='resume'):
        read = documents.read_attachment(self.ctx, self.attachment.pk)
        return documents.propose_document(self.ctx, read['document_id'], kind,
            [{'field': 'name', 'value': 'Asha Raman', 'evidence_quote': 'Asha Raman'},
             {'field': 'gpa', 'value': 3.8, 'evidence_quote': 'GPA 3.8/4.0'}], 'Academic profile')

    def test_resume_confirmation_updates_main_profile_and_preserves_other_sources(self):
        StudentProfile.objects.filter(pk=self.student.pk).update(github_assessment={'repo': 'keep'}, linkedin_profile={'name': 'LinkedIn name'})
        proposal = self.document_proposal()
        self.student.refresh_from_db()
        self.assertEqual(self.student.name, 'Asha')
        self.assertFalse(ResumeUpload.objects.exists())
        changes.resolve(self.context('yes', 'two'), proposal['id'], 'approve', 'yes')
        self.student.refresh_from_db()
        self.assertEqual((self.student.name, self.student.gpa), ('Asha Raman', 3.8))
        self.assertEqual(self.student.github_assessment, {'repo': 'keep'})
        self.assertEqual(self.student.linkedin_profile, {'name': 'LinkedIn name'})
        self.assertEqual(ResumeUpload.objects.count(), 1)

    def test_linkedin_confirmation_only_updates_linkedin(self):
        proposal = self.document_proposal('linkedin')
        changes.resolve(self.context('yes', 'two'), proposal['id'], 'approve', 'yes')
        self.student.refresh_from_db()
        self.assertEqual((self.student.name, self.student.gpa), ('Asha', 3.4))
        self.assertEqual(self.student.linkedin_profile['gpa'], 3.8)
        self.assertEqual(LinkedInAnalysis.objects.count(), 1)
        self.assertEqual(ResumeUpload.objects.count(), 0)

    def test_declining_document_leaves_all_sources_and_assumptions_unchanged(self):
        proposal = self.document_proposal()
        ctx = self.context('no', 'two')
        changes.resolve(ctx, proposal['id'], 'reject', 'no')
        self.assertFalse(ResumeUpload.objects.exists())
        self.assertEqual(documents.confirmed_evidence(str(self.student.uuid)), {})
        self.assertFalse(changes.conversation_state(ctx)['conversation_assumptions'])

    def test_quotes_and_attachment_ownership_checked(self):
        read = documents.read_attachment(self.ctx, self.attachment.pk)
        with self.assertRaisesMessage(ValueError, 'verbatim'):
            documents.propose_document(self.ctx, read['document_id'], 'resume',
                [{'field': 'name', 'value': 'Invented', 'evidence_quote': 'Missing quote'}], 'Summary')
        ctx = {**self.ctx, 'canonical_student_id': str(uuid.uuid4())}
        with self.assertRaisesMessage(ValueError, 'not found'):
            documents.read_attachment(ctx, self.attachment.pk)

    def test_unfinished_extraction_survives_a_new_turn_and_clear_cancels_it(self):
        read = documents.read_attachment(self.ctx, self.attachment.pk)
        restored = documents.unfinished_documents(str(self.student.uuid))
        self.assertTrue(restored[str(read['document_id'])]['needs_proposal'])
        self.assertIn('Asha Raman', restored[str(read['document_id'])]['text'])
        changes.clear_conversation(student_id=str(self.student.uuid))
        self.assertEqual(documents.unfinished_documents(str(self.student.uuid)), {})

    def test_review_only_does_not_require_update_proposal(self):
        documents.read_attachment(self.ctx, self.attachment.pk, purpose='review')
        self.assertEqual(documents.unfinished_documents(str(self.student.uuid)), {})

    @override_settings(AGENT_DISTRIBUTED_LIMITS=False)
    def test_native_graph_requires_document_proposal_before_reply(self):
        from pure_multi_agent.student_graph import build_student_agent
        read = documents.read_attachment(self.ctx, self.attachment.pk)
        self.ctx['documents_read'] = documents.unfinished_documents(str(self.student.uuid))
        steps = []
        def model(messages, tools, **kwargs):
            steps.append(kwargs.get('require_tools', False))
            if len(steps) == 1:
                self.assertEqual({t.name for t in tools}, {'read_student_document', 'propose_document_update', 'finish_document_review', 'request_document_clarification', 'discard_document_draft'})
                return AIMessage(content='', tool_calls=[{'id': 'doc', 'name': 'propose_document_update', 'args': {
                    'document_id': read['document_id'], 'source_type': 'resume', 'summary': 'Resume facts',
                    'facts': [{'field': 'name', 'value': 'Asha Raman', 'evidence_quote': 'Asha Raman'}]}}])
            return AIMessage(content='Please confirm the resume update.')
        with mock.patch('pure_multi_agent.model_router.invoke', side_effect=model):
            agent = build_student_agent(self.ctx, 'Assist the student.', InMemorySaver())
            agent.invoke({'messages': [HumanMessage(content='Continue my resume update')]}, {'configurable': {'thread_id': 'doc-test'}})
        self.assertEqual(steps, [True, False])
        self.assertTrue(AgentChangeProposal.objects.filter(kind='student_document', status='pending').exists())
        self.student.refresh_from_db()
        self.assertEqual(self.student.name, 'Asha')

    def test_pdf_text_is_available_to_qwen_without_vision(self):
        from reportlab.pdfgen import canvas
        path = Path(self.temp.name) / 'resume.pdf'
        pdf = canvas.Canvas(str(path)); pdf.drawString(40, 750, 'Asha Raman GPA 3.8/4.0'); pdf.save()
        self.attachment.file_path, self.attachment.content_type = str(path), 'application/pdf'
        self.attachment.save()
        result = documents.read_attachment(self.ctx, self.attachment.pk)
        self.assertIn('Asha Raman', result['text'])
        self.assertNotIn('_media_blocks', result)

    def test_docx_table_text_is_read(self):
        from docx import Document
        path = Path(self.temp.name) / 'resume.docx'
        doc = Document(); doc.add_paragraph('Asha Raman'); doc.add_table(rows=1, cols=1).cell(0, 0).text = 'GPA 3.8'; doc.save(path)
        self.attachment.file_path, self.attachment.content_type = str(path), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        self.attachment.save()
        self.assertIn('GPA 3.8', documents.read_attachment(self.ctx, self.attachment.pk)['text'])


@override_settings(UNIVERSITY_VECTOR_SEARCH=False, AGENT_DISTRIBUTED_LIMITS=False)
class OfficerTests(TestCase):
    def setUp(self):
        self.uni = University.objects.create(name='Test University')
        self.actor = User.objects.create_user(username='officer')
        Account.objects.create(user=self.actor, role='university', university=self.uni)
        self.ctx = {'university_id': str(self.uni.uuid), 'actor_id': self.actor.pk, 'turn_id': 'one', 'current_message': 'Add a scholarship'}

    def propose(self):
        return changes.propose_university(self.ctx, 'university_knowledge', {'topic': 'Merit scholarship', 'content': '$2000 for eligible students', 'group': 'money', 'details': {'category': 'scholarship', 'applies_to': 'all graduate applicants', 'effective_period': 'Fall 2027', 'amount': 2000, 'currency': 'USD', 'amount_basis': 'per academic year', 'eligibility': 'CGPA 3.5/4.0 or above', 'application_process': 'Apply online', 'deadline': '2027-01-31'}})

    def confirm(self, proposal, decision='approve'):
        return changes.resolve({**self.ctx, 'turn_id': 'two', 'current_message': 'yes'}, proposal['id'], decision, 'yes')

    def test_create_policy_waits_then_indexes_and_is_idempotent(self):
        proposal = self.propose()
        self.assertFalse(UniversityKnowledgeEntry.objects.exists())
        with self.assertRaises(ValueError):
            changes.resolve(self.ctx, proposal['id'], 'approve', self.ctx['current_message'])
        for _ in range(2):
            self.assertEqual(self.confirm(proposal)['status'], 'applied')
        self.assertEqual(UniversityKnowledgeEntry.objects.count(), 1)
        self.assertTrue(KnowledgeIndexWork.objects.filter(university_id=str(self.uni.uuid)).exists())
        entry = UniversityKnowledgeEntry.objects.get()
        self.assertEqual((entry.source_type, entry.group.slug), ('officer', 'money'))

    def test_policy_update_stale_guard(self):
        row = UniversityKnowledgeEntry.objects.create(university_id=str(self.uni.uuid), topic='Scholarship', content='Old')
        proposal = changes.propose_university(self.ctx, 'university_knowledge', {'topic': 'Scholarship', 'content': 'New', 'group': 'money', 'details': {'category': 'scholarship', 'applies_to': 'all students', 'effective_period': '2027', 'coverage': 'full tuition', 'eligibility': 'admitted students', 'application_process': 'automatic consideration', 'deadline': 'no deadline'}}, row.pk)
        row.content = 'Changed elsewhere'; row.save()
        self.assertEqual(self.confirm(proposal)['status'], 'stale')
        row.refresh_from_db()
        self.assertEqual(row.content, 'Changed elsewhere')

    def test_actor_and_university_cannot_cross_boundaries(self):
        proposal = self.propose()
        second = User.objects.create_user(username='second')
        Account.objects.create(user=second, role='university', university=self.uni)
        with self.assertRaises(ValueError):
            changes.resolve({**self.ctx, 'actor_id': second.pk, 'turn_id': 'two', 'current_message': 'yes'}, proposal['id'], 'approve', 'yes')
        self.actor.is_active = False; self.actor.save()
        with self.assertRaises(ValueError):
            self.confirm(proposal)

    def test_requirement_tool_preserves_unrelated_criteria(self):
        from pure_multi_agent.tools.officer_tools import build_tools
        self.uni.eligibility_criteria = [{'criterion': 'Degree', 'detail': 'Bachelor degree'}]; self.uni.save()
        tool = next(t for t in build_tools(self.ctx) if t.name == 'propose_admission_requirement')
        proposal = tool.invoke({'operation': 'add', 'requirement': {'criterion': 'IELTS', 'detail': 'At least 6.5 out of 9', 'category': 'test_score', 'applies_to': 'all applicants', 'test_name': 'IELTS Academic', 'minimum': 6.5, 'maximum': 9, 'scale_maximum': 9}})
        self.assertEqual(self.confirm(proposal)['status'], 'applied')
        self.uni.refresh_from_db()
        self.assertEqual(len(self.uni.eligibility_criteria), 2)
        self.assertEqual(self.uni.eligibility_criteria[0]['criterion'], 'Degree')

    def test_student_cards_only_for_own_interested_students(self):
        from pure_multi_agent.tools.officer_tools import build_tools
        student = StudentProfile.objects.create(name='Asha')
        other = StudentProfile.objects.create(name='Private')
        UniversityInterestEvent.objects.create(student=student, university_id=str(self.uni.uuid), source='searched')
        tool = next(t for t in build_tools(self.ctx) if t.name == 'interested_student_detail')
        allowed = tool.invoke({'student_id': str(student.uuid)})
        self.assertEqual(allowed['student_cards'][0]['name'], 'Asha')
        self.assertIn(f'/profiles/{student.uuid}', allowed['student_cards'][0]['profile_path'])
        self.assertIn('error', tool.invoke({'student_id': str(other.uuid)}))
        self.assertNotIn(str(other.uuid), self.ctx['student_cards'])

    def test_native_graph_calls_tool_then_resumes_with_confirmation(self):
        from pure_multi_agent.officer_graph import run_turn
        saver = InMemorySaver()
        calls = []
        def model(messages, tools, **kwargs):
            calls.append([t.name for t in tools])
            last = messages[-1]
            if last.type == 'tool':
                return AIMessage(content='Please confirm.' if 'pending' in last.content else 'Saved.')
            if last.content == 'yes':
                proposal = AgentChangeProposal.objects.get(status='pending')
                name, args = 'resolve_university_change', {'proposal_id': str(proposal.pk), 'decision': 'approve', 'confirmation_message': 'yes'}
            else:
                name, args = 'propose_knowledge_change', {'topic': 'Merit', 'content': 'Applicants may request advice', 'group': 'money', 'details': {'category': 'general', 'applies_to': 'all applicants', 'effective_period': 'ongoing'}}
            return AIMessage(content='', tool_calls=[{'name': name, 'args': args, 'id': str(uuid.uuid4())}])
        with mock.patch('pure_multi_agent.model_router.invoke', side_effect=model):
            result = run_turn(str(self.uni.uuid), self.actor.pk, 'Add a $2000 award', checkpointer=saver)
            self.assertEqual(result['change_proposals'][0]['status'], 'pending')
            result = run_turn(str(self.uni.uuid), self.actor.pk, 'yes', checkpointer=saver)
        self.assertEqual(result['change_proposals'][0]['status'], 'applied')
        self.assertTrue(all('interested_student_detail' in names for names in calls))
        self.assertEqual(UniversityKnowledgeEntry.objects.count(), 1)

    def test_graph_capacity_resumes_without_replaying_policy_creation(self):
        from pure_multi_agent.officer_graph import run_turn
        from pure_multi_agent.capacity import ResumeTurnLater
        from github_profiles.scheduling import CapacityBusy
        saver = InMemorySaver()
        proposal_call = AIMessage(content='', tool_calls=[{'name': 'propose_knowledge_change', 'args': {'topic': 'Merit', 'content': 'Applicants may request advice', 'group': 'money', 'details': {'category': 'general', 'applies_to': 'all applicants', 'effective_period': 'ongoing'}}, 'id': 'call'}])
        with mock.patch('pure_multi_agent.model_router.invoke', side_effect=[proposal_call, CapacityBusy('busy', 5)]):
            with self.assertRaises(ResumeTurnLater) as raised:
                run_turn(str(self.uni.uuid), self.actor.pk, 'Add award', checkpointer=saver)
        with mock.patch('pure_multi_agent.model_router.invoke', return_value=AIMessage(content='Confirm?')):
            result = run_turn(str(self.uni.uuid), self.actor.pk, 'Add award', checkpointer=saver, resume_state=raised.exception.state)
        self.assertEqual(AgentChangeProposal.objects.count(), 1)
        self.assertEqual(result['change_proposals'][0]['status'], 'pending')
